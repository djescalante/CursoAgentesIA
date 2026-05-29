"""
Agente: Procesador de Indicadores ANS por Operador
Versión 1.2 - Procesamiento de incidentes y generación de reportes con soporte de fórmulas y cálculo de cumplimiento
"""

import pandas as pd
import numpy as np
import re
import calendar
from pathlib import Path
from datetime import datetime, timedelta
from openpyxl import load_workbook
from openpyxl.styles import Font, Border, Side, Alignment, PatternFill
import shutil


class ProcesadorIndicadoresANS:
    """Clase principal que orquesta el procesamiento de indicadores ANS"""
    
    def __init__(self, verbose=True):
        self.verbose = verbose
    
    def log(self, mensaje):
        """Imprime mensajes si verbose está activado"""
        if self.verbose:
            print(mensaje)
    
    # ==================== SKILL 1: Leer Excel ====================
    def load_excel_data(self, file_path: str, sheet_name=0) -> pd.DataFrame:
        """Carga datos desde un archivo Excel"""
        try:
            path = Path(file_path)
            if not path.exists():
                self.log(f"[ERROR] Archivo {file_path} no encontrado")
                return None
            
            # Cargar con pandas
            df = pd.read_excel(file_path, sheet_name=sheet_name)
            
            self.log(f"[OK] Datos cargados exitosamente")
            self.log(f"  - Filas: {df.shape[0]}")
            self.log(f"  - Columnas: {df.shape[1]}")
            
            return df
        
        except Exception as e:
            self.log(f"[ERROR] Error al cargar archivo: {str(e)}")
            return None
    
    # ==================== SKILL 2: Normalizar Columnas ====================
    def normalize_columns(self, df: pd.DataFrame) -> pd.DataFrame:
        """Normaliza los nombres de las columnas"""
        df = df.copy()
        
        cambios = {}
        for col in df.columns:
            col_normalizada = col.lower()
            col_normalizada = col_normalizada.replace(' ', '_')
            col_normalizada = re.sub(r'[áéíóú]', lambda m: {
                'á': 'a', 'é': 'e', 'í': 'i', 'ó': 'o', 'ú': 'u'
            }[m.group()], col_normalizada)
            col_normalizada = re.sub(r'[^a-z0-9_]', '', col_normalizada)
            
            if col != col_normalizada:
                cambios[col] = col_normalizada
        
        df.rename(columns=cambios, inplace=True)
        
        if cambios:
            self.log(f"[OK] Columnas normalizadas: {len(cambios)} cambios")
        
        return df
    
    # ==================== SKILL 3: Mapear Variantes ====================
    def map_column_variants(self, df: pd.DataFrame, mapeo: dict = None) -> pd.DataFrame:
        """Mapea columnas con nombres variantes a nombres estándar"""
        if mapeo is None:
            mapeo = {
                'operador': ['reconocidopornombre', 'operador_reconocimiento', 'reconocido_por_nombre'],
                'oportunidad': ['cumplimiento_ans', 'oportunidad'],
                'prioridad': ['prioridad', 'prioridades'],
                'fecha': ['fechacreacion', 'fecha_creacion'],
                'ans_val': ['ans']
            }
        
        df_procesado = df.copy()
        
        for nombre_estandar, variantes in mapeo.items():
            for col in df_procesado.columns:
                col_norm = col.lower()
                if col_norm in variantes:
                    if col != nombre_estandar:
                        df_procesado.rename(columns={col: nombre_estandar}, inplace=True)
                        self.log(f"[OK] Mapeada columna: {col} -> {nombre_estandar}")
                    break
        
        return df_procesado
    
    # ==================== SKILL 4: Validar Datos y Mapear Valores ====================
    def validate_data(self, df: pd.DataFrame, required_columns: list) -> tuple:
        """Valida y limpia los datos del DataFrame, mapeando prioridades y calculando oportunidades si faltan"""
        df_validado = df.copy()
        
        # 1. Verificar columnas requeridas básicas
        for col in ['operador', 'prioridad', 'fecha']:
            if col not in df_validado.columns:
                self.log(f"[ERROR] Columna requerida faltante: {col}")
                return (False, f"Columna faltante: {col}", None)
                
        # 2. Eliminar filas vacías en columnas clave
        df_validado = df_validado.dropna(subset=['operador', 'prioridad', 'fecha'])
            
        # 3. Normalizar y mapear prioridades (manejar 1, 2, 3 y strings como 'alta', 'media', 'baja')
        priority_mapping = {
            1: 'Alta', 2: 'Media', 3: 'Baja',
            '1': 'Alta', '2': 'Media', '3': 'Baja',
            'alta': 'Alta', 'media': 'Media', 'baja': 'Baja',
            'Alta': 'Alta', 'Media': 'Media', 'Baja': 'Baja'
        }
        df_validado['prioridad'] = df_validado['prioridad'].map(priority_mapping).fillna(df_validado['prioridad'].astype(str).str.capitalize())
        df_validado = df_validado[df_validado['prioridad'].isin(['Alta', 'Media', 'Baja'])]
        
        # 4. Normalizar oportunidad si existe
        opportunity_mapping = {
            'oportuno': 'Oportuno',
            'no oportuno': 'No Oportuno',
            'no_oportuno': 'No Oportuno',
            'Oportuno': 'Oportuno',
            'No Oportuno': 'No Oportuno'
        }
        
        if 'oportunidad' in df_validado.columns:
            df_validado['oportunidad'] = df_validado['oportunidad'].astype(str).str.strip().map(opportunity_mapping)
        else:
            df_validado['oportunidad'] = np.nan
            
        # 5. Rellenar oportunidad usando reglas de ANS si es nula pero ans_val está disponible
        if 'ans_val' in df_validado.columns:
            def calcular_oportunidad(row):
                opt = row['oportunidad']
                if pd.notna(opt):
                    return opt
                
                ans = row['ans_val']
                prio = row['prioridad']
                if pd.isna(ans) or pd.isna(prio):
                    return None
                    
                try:
                    ans_num = float(ans)
                    if prio == 'Alta':
                        return 'Oportuno' if ans_num <= 60 else 'No Oportuno'
                    elif prio == 'Media':
                        return 'Oportuno' if ans_num <= 300 else 'No Oportuno'
                    elif prio == 'Baja':
                        return 'Oportuno' if ans_num <= 600 else 'No Oportuno'
                except:
                    pass
                return None
                
            df_validado['oportunidad'] = df_validado.apply(calcular_oportunidad, axis=1)
            
        # Eliminar filas donde oportunidad sigue siendo nula
        df_validado = df_validado.dropna(subset=['oportunidad'])
        df_validado = df_validado[df_validado['oportunidad'].isin(['Oportuno', 'No Oportuno'])]
        
        self.log(f"[OK] Validacion y mapeo completados. Filas validas: {len(df_validado)}")
        return (True, "Validacion exitosa", df_validado)
    
    # ==================== SKILL 5: Crear Períodos ====================
    def create_periods(self, df: pd.DataFrame, fecha_column: str = 'fecha') -> pd.DataFrame:
        """Crea períodos de ~8 días de manera dinámica y consistente con Enero"""
        df_procesado = df.copy()
        
        df_procesado[fecha_column] = pd.to_datetime(df_procesado[fecha_column], errors='coerce')
        
        meses_esp = {
            1: 'Enero', 2: 'Febrero', 3: 'Marzo', 4: 'Abril', 5: 'Mayo', 6: 'Junio',
            7: 'Julio', 8: 'Agosto', 9: 'Septiembre', 10: 'Octubre', 11: 'Noviembre', 12: 'Diciembre'
        }
        
        def mapear_fecha_a_periodo(fecha):
            if pd.isna(fecha):
                return None
            
            day = fecha.day
            month_num = fecha.month
            month_name = meses_esp.get(month_num, fecha.strftime('%B'))
            
            if 1 <= day <= 8:
                return f"{month_name} 1-8"
            elif 9 <= day <= 15:
                return f"{month_name} 9-15"
            elif 16 <= day <= 22:
                return f"{month_name} 16-22"
            else:
                # Buscar el día máximo para este período de fin de mes
                month_df = df_procesado[df_procesado[fecha_column].dt.month == month_num]
                max_day_in_data = month_df[fecha_column].dt.day.max()
                if pd.isna(max_day_in_data) or max_day_in_data < 23:
                    _, max_day_in_data = calendar.monthrange(fecha.year, month_num)
                
                # Respetar espaciado de Enero si es el caso
                if month_num == 1 and max_day_in_data == 29:
                    return "Enero 23- 29"
                return f"{month_name} 23-{max_day_in_data}"
        
        df_procesado['periodo'] = df_procesado[fecha_column].apply(mapear_fecha_a_periodo)
        
        self.log("[OK] Periodos generados:")
        unique_periods = sorted([p for p in df_procesado['periodo'].dropna().unique()])
        for p in unique_periods:
            count = len(df_procesado[df_procesado['periodo'] == p])
            self.log(f"  - {p}: {count} incidentes")
            
        return df_procesado
    
    # ==================== SKILL 6: Agrupar por Operador y Período ====================
    def group_by_operator_period(self, df: pd.DataFrame, operador_column: str = 'operador',
                                periodo_column: str = 'periodo') -> dict:
        """Agrupa incidentes por operador y período"""
        grupos = {}
        
        for periodo, df_periodo in df.groupby(periodo_column):
            grupos[periodo] = {}
            for operador, df_operador in df_periodo.groupby(operador_column):
                grupos[periodo][operador] = df_operador
        
        return grupos
    
    # ==================== SKILL 7: Calcular Indicadores ====================
    def calculate_indicators(self, grupos: dict, oportunidad_column: str = 'oportunidad',
                           prioridad_column: str = 'prioridad') -> pd.DataFrame:
        """Calcula agregaciones básicas de oportunidad por operador y prioridad"""
        datos = []
        
        for periodo, ops in grupos.items():
            for operador, df_operador in ops.items():
                row = {
                    'periodo': periodo,
                    'operador': operador,
                    'Alta_NoOport': 0, 'Alta_Total': 0,
                    'Baja_NoOport': 0, 'Baja_Total': 0,
                    'Media_NoOport': 0, 'Media_Total': 0
                }
                
                for prioridad in ['Alta', 'Media', 'Baja']:
                    df_p = df_operador[df_operador[prioridad_column] == prioridad]
                    total = len(df_p)
                    no_oport = len(df_p[df_p[oportunidad_column] == 'No Oportuno'])
                    
                    if prioridad == 'Alta':
                        row['Alta_NoOport'] = no_oport
                        row['Alta_Total'] = total
                    elif prioridad == 'Baja':
                        row['Baja_NoOport'] = no_oport
                        row['Baja_Total'] = total
                    elif prioridad == 'Media':
                        row['Media_NoOport'] = no_oport
                        row['Media_Total'] = total
                
                datos.append(row)
                
        return pd.DataFrame(datos)
    
    # ==================== SKILL 9: Exportar a Excel con Fórmulas y Formatos ====================
    def export_indicators_to_excel(self, input_file_path: str, output_file_path: str,
                                  df_indicadores: pd.DataFrame, df_original: pd.DataFrame,
                                  sheet_name: str = 'indicador alarmas') -> bool:
        """Exporta los indicadores calculados al archivo Excel aplicando fórmulas y estilos exactos"""
        try:
            Path(output_file_path).parent.mkdir(parents=True, exist_ok=True)
            
            # Copiar archivo original
            shutil.copy(input_file_path, output_file_path)
            
            # Cargar workbook
            wb = load_workbook(output_file_path)
            
            # Renombrar Hoja1 (o la de origen) a "alarmas" para cumplir con el esquema
            sheet_origen = wb.sheetnames[0]
            if sheet_origen != 'alarmas':
                ws_orig = wb[sheet_origen]
                ws_orig.title = 'alarmas'
                self.log(f"[OK] Renombrada hoja de datos originales '{sheet_origen}' -> 'alarmas'")
            
            # Crear o limpiar hoja de indicadores
            if sheet_name in wb.sheetnames:
                del wb[sheet_name]
            ws = wb.create_sheet(sheet_name, 0)
            
            # Estilos
            title_font = Font(bold=True, size=11, name="Calibri")
            header_font = Font(bold=True, size=11, name="Calibri")
            data_font = Font(size=11, name="Calibri")
            header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
            total_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
            
            thin_border = Border(
                left=Side(style='thin', color='D9D9D9'),
                right=Side(style='thin', color='D9D9D9'),
                top=Side(style='thin', color='D9D9D9'),
                bottom=Side(style='thin', color='D9D9D9')
            )
            
            headers = [
                'Operador', 'No oport.', 'Alta ', 'Efect.', 
                'No oport.2', 'Baja', 'Efect.3', 
                'No oport.4', 'Media', 'Efect.5', 
                'Total Incidentes gestionados'
            ]
            
            fila_actual = 1
            
            # Agrupar por período
            for periodo in sorted(df_indicadores['periodo'].unique()):
                df_p = df_indicadores[df_indicadores['periodo'] == periodo].copy()
                df_p = df_p.sort_values(by='operador') # Ordenar alfabéticamente
                
                # Título del Período
                ws.cell(row=fila_actual, column=1, value=periodo).font = title_font
                fila_actual += 1
                
                # Escribir Headers
                for col_idx, h in enumerate(headers, 1):
                    cell = ws.cell(row=fila_actual, column=col_idx, value=h)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.border = thin_border
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                
                ws.row_dimensions[fila_actual].height = 28
                fila_header = fila_actual
                fila_actual += 1
                
                start_data_row = fila_actual
                
                # Escribir Datos de Operadores
                for _, row in df_p.iterrows():
                    ws.cell(row=fila_actual, column=1, value=row['operador']).font = data_font
                    
                    # Alta
                    ws.cell(row=fila_actual, column=2, value=row['Alta_NoOport']).font = data_font
                    ws.cell(row=fila_actual, column=3, value=row['Alta_Total']).font = data_font
                    cell_ef_alta = ws.cell(row=fila_actual, column=4, value=f"=(C{fila_actual}-B{fila_actual})/C{fila_actual}")
                    cell_ef_alta.font = data_font
                    cell_ef_alta.number_format = '0.0%'
                    
                    # Baja
                    ws.cell(row=fila_actual, column=5, value=row['Baja_NoOport']).font = data_font
                    ws.cell(row=fila_actual, column=6, value=row['Baja_Total']).font = data_font
                    cell_ef_baja = ws.cell(row=fila_actual, column=7, value=f"=IFERROR((F{fila_actual}-E{fila_actual})/F{fila_actual},1)")
                    cell_ef_baja.font = data_font
                    cell_ef_baja.number_format = '0.0%'
                    
                    # Media (No oport.4 es None si es 0, para consistencia con enero)
                    no_oport_media = row['Media_NoOport'] if row['Media_NoOport'] > 0 else None
                    ws.cell(row=fila_actual, column=8, value=no_oport_media).font = data_font
                    ws.cell(row=fila_actual, column=9, value=row['Media_Total']).font = data_font
                    cell_ef_media = ws.cell(row=fila_actual, column=10, value=f"=IFERROR((I{fila_actual}-H{fila_actual})/I{fila_actual},1)")
                    cell_ef_media.font = data_font
                    cell_ef_media.number_format = '0.0%'
                    
                    # Total
                    cell_total = ws.cell(row=fila_actual, column=11, value=f"=SUM(C{fila_actual},F{fila_actual},I{fila_actual})")
                    cell_total.font = data_font
                    
                    # Alinear al centro columnas numéricas
                    for col_c in range(1, 12):
                        cell_item = ws.cell(row=fila_actual, column=col_c)
                        cell_item.border = thin_border
                        if col_c > 1:
                            cell_item.alignment = Alignment(horizontal='center')
                    
                    fila_actual += 1
                
                end_data_row = fila_actual - 1
                
                # Fila de Total General
                ws.cell(row=fila_actual, column=1, value="Total general").font = Font(bold=True, size=11)
                
                # Sumas de No Oportunos y Totales por prioridad
                ws.cell(row=fila_actual, column=2, value=f"=SUM(B{start_data_row}:B{end_data_row})")
                ws.cell(row=fila_actual, column=3, value=f"=SUM(C{start_data_row}:C{end_data_row})")
                cell_tot_ef_alta = ws.cell(row=fila_actual, column=4, value=f"=(C{fila_actual}-B{fila_actual})/C{fila_actual}")
                cell_tot_ef_alta.number_format = '0.0%'
                
                ws.cell(row=fila_actual, column=5, value=f"=SUM(E{start_data_row}:E{end_data_row})")
                ws.cell(row=fila_actual, column=6, value=f"=SUM(F{start_data_row}:F{end_data_row})")
                cell_tot_ef_baja = ws.cell(row=fila_actual, column=7, value=f"=(F{fila_actual}-E{fila_actual})/F{fila_actual}")
                cell_tot_ef_baja.number_format = '0.0%'
                
                ws.cell(row=fila_actual, column=8, value=f"=SUM(H{start_data_row}:H{end_data_row})")
                ws.cell(row=fila_actual, column=9, value=f"=SUM(I{start_data_row}:I{end_data_row})")
                cell_tot_ef_media = ws.cell(row=fila_actual, column=10, value=f"=(I{fila_actual}-H{fila_actual})/I{fila_actual}")
                cell_tot_ef_media.number_format = '0.0%'
                
                ws.cell(row=fila_actual, column=11, value=f"=SUM(K{start_data_row}:K{end_data_row})")
                
                # Aplicar estilos a Fila de Totales
                for col_c in range(1, 12):
                    cell_tot = ws.cell(row=fila_actual, column=col_c)
                    cell_tot.font = Font(bold=True, size=11, name="Calibri")
                    cell_tot.fill = total_fill
                    cell_tot.border = thin_border
                    if col_c > 1:
                        cell_tot.alignment = Alignment(horizontal='center')
                
                fila_actual += 3  # Dos filas vacías de espacio
                
            # Ajustar anchos de columnas
            for col in range(1, 12):
                col_letter = ws.cell(row=1, column=col).column_letter
                max_len = 0
                for row in range(1, fila_actual):
                    cell_val = ws.cell(row=row, column=col).value
                    if cell_val is not None:
                        max_len = max(max_len, len(str(cell_val)))
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)
            
            # Guardar workbook
            wb.save(output_file_path)
            self.log("[OK] Guardado final del archivo Excel completado con exito")
            return True
            
        except Exception as e:
            self.log(f"[ERROR] Error al exportar indicadores a Excel: {str(e)}")
            import traceback
            self.log(traceback.format_exc())
            return False
            
    # ==================== SKILL 10: Crear Reporte ====================
    def create_processing_report(self, df_original: pd.DataFrame, df_indicadores: pd.DataFrame, 
                                 output_file_path: str) -> str:
        """Crea un reporte consolidado del procesamiento"""
        reporte = []
        reporte.append("=" * 70)
        reporte.append("REPORTE DE PROCESAMIENTO - INDICADORES ANS MAYO 2026")
        reporte.append("=" * 70)
        reporte.append(f"Fecha y Hora: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        reporte.append("")
        reporte.append(f"Total de incidentes procesados: {len(df_original):,}")
        reporte.append(f"Operadores unicos procesados: {df_indicadores['operador'].nunique()}")
        reporte.append(f"Periodos procesados: {df_indicadores['periodo'].nunique()}")
        reporte.append("")
        reporte.append("RESUMEN DE PERIODOS:")
        for p in sorted(df_indicadores['periodo'].unique()):
            df_p = df_indicadores[df_indicadores['periodo'] == p]
            ops_count = df_p['operador'].nunique()
            reporte.append(f"  - {p}: {ops_count} operadores activos")
        reporte.append("")
        reporte.append(f"Archivo de salida guardado en:\n  {output_file_path}")
        reporte.append("=" * 70)
        return "\n".join(reporte)
        
    def ejecutar(self, input_file_path: str, output_file_path: str) -> bool:
        """Ejecuta el pipeline completo del agente"""
        self.log(f"Iniciando procesamiento de archivo: {input_file_path}")
        
        # 1. Cargar datos
        df = self.load_excel_data(input_file_path)
        if df is None:
            return False
            
        # 2. Normalizar columnas
        df = self.normalize_columns(df)
        
        # 3. Mapear variantes de nombres de columnas
        df = self.map_column_variants(df)
        
        # 4. Validar y limpiar datos
        required_cols = ['operador', 'prioridad', 'fecha']
        es_valido, msg, df_limpio = self.validate_data(df, required_cols)
        if not es_valido:
            self.log(f"[ERROR] Error de validacion: {msg}")
            return False
            
        # 5. Crear periodos de fecha
        df_con_periodos = self.create_periods(df_limpio, 'fecha')
        
        # 6. Agrupar por operador y periodo
        grupos = self.group_by_operator_period(df_con_periodos)
        
        # 7. Calcular indicadores por prioridad
        df_indicadores = self.calculate_indicators(grupos)
        
        # 8. Exportar a Excel con formulas y formatos
        exito = self.export_indicators_to_excel(input_file_path, output_file_path, df_indicadores, df_con_periodos)
        if not exito:
            return False
            
        # 9. Mostrar reporte
        reporte = self.create_processing_report(df_con_periodos, df_indicadores, output_file_path)
        print("\n" + reporte)
        
        return True


if __name__ == "__main__":
    # Detectar directorio relativo para ejecucion
    script_dir = Path(__file__).resolve().parent.parent
    input_file = script_dir / "Incidentes_sin_procesar" / "Incidentes mayo 01-27-2026.xlsx"
    output_file = script_dir / "Incidentes_ANS_Procesado" / "Incidentes_mayo_2026_procesado.xlsx"
    
    procesador = ProcesadorIndicadoresANS(verbose=True)
    procesador.ejecutar(str(input_file), str(output_file))
